from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from openpyxl import load_workbook


BASE = Path(
    "/Users/hendrikribeiro/Library/Mobile Documents/com~apple~CloudDocs/"
    "CODE HUAMBO/untitled folder 2/fotos 22 de maio"
)

FILES = {
    "monthly_report": BASE / "HMRL_Relatorio Mensal HSE_36_2026.05_Rev.01.docx",
    "checklist": BASE / "HMRL_Checklist RM_36_2026-05.xlsm",
    "incident_board": BASE / "HMRL_Quadro de Sinistralidade_RM 36_2026.05.xlsx",
}

OUT = Path("extracted-luena")


def norm(value):
    if value is None:
        return ""
    text = str(value).replace("\xa0", " ").strip()
    return " ".join(text.split())


def extract_docx(path: Path) -> dict:
    document = Document(path)
    paragraphs = [norm(p.text) for p in document.paragraphs if norm(p.text)]
    tables = []
    for table_index, table in enumerate(document.tables, start=1):
        rows = []
        for row in table.rows:
            values = [norm(cell.text) for cell in row.cells]
            if any(values):
                rows.append(values)
        tables.append({"table": table_index, "rows": rows})
    return {"path": str(path), "paragraphs": paragraphs, "tables": tables}


def extract_xlsx(path: Path) -> dict:
    workbook = load_workbook(path, data_only=False, keep_vba=path.suffix.lower() == ".xlsm")
    sheets = []
    for sheet in workbook.worksheets:
        rows = []
        for row in sheet.iter_rows():
            values = [norm(cell.value) for cell in row]
            if any(values):
                rows.append({"row": row[0].row, "values": values})
        sheets.append(
            {
                "title": sheet.title,
                "max_row": sheet.max_row,
                "max_column": sheet.max_column,
                "rows": rows,
            }
        )
    return {"path": str(path), "sheets": sheets}


def write_doc_markdown(name: str, data: dict) -> None:
    lines = [f"# {name}", "", f"Source: `{data['path']}`", ""]
    if data["paragraphs"]:
        lines.extend(["## Paragraphs", ""])
        for i, paragraph in enumerate(data["paragraphs"], start=1):
            lines.append(f"{i}. {paragraph}")
        lines.append("")
    for table in data["tables"]:
        lines.extend([f"## Table {table['table']}", ""])
        for row in table["rows"]:
            lines.append(" | ".join(row))
        lines.append("")
    (OUT / f"{name}.md").write_text("\n".join(lines), encoding="utf-8")


def write_xlsx_markdown(name: str, data: dict) -> None:
    lines = [f"# {name}", "", f"Source: `{data['path']}`", ""]
    for sheet in data["sheets"]:
        lines.extend([f"## Sheet: {sheet['title']} ({sheet['max_row']} x {sheet['max_column']})", ""])
        for row in sheet["rows"]:
            lines.append(f"R{row['row']}: " + " | ".join(row["values"]))
        lines.append("")
    (OUT / f"{name}.md").write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    OUT.mkdir(exist_ok=True)
    manifest = {}
    for name, path in FILES.items():
        if not path.exists():
            manifest[name] = {"missing": str(path)}
            continue
        if path.suffix.lower() == ".docx":
            data = extract_docx(path)
            write_doc_markdown(name, data)
            manifest[name] = {
                "paragraphs": len(data["paragraphs"]),
                "tables": len(data["tables"]),
                "table_rows": [len(table["rows"]) for table in data["tables"]],
            }
        else:
            data = extract_xlsx(path)
            write_xlsx_markdown(name, data)
            manifest[name] = {
                "sheets": [
                    {
                        "title": sheet["title"],
                        "max_row": sheet["max_row"],
                        "max_column": sheet["max_column"],
                        "non_empty_rows": len(sheet["rows"]),
                    }
                    for sheet in data["sheets"]
                ]
            }
    print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
