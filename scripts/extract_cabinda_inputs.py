from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from openpyxl import load_workbook

try:
    from pypdf import PdfReader
except Exception:  # pragma: no cover - optional extraction only
    PdfReader = None


BASE = Path(
    "/Users/hendrikribeiro/Library/Mobile Documents/com~apple~CloudDocs/"
    "CODE HUAMBO/untitled folder 2/fotos 22 de maio"
)

FILES = {
    "monthly_report": BASE / "HMRC_Relatorio Mensal HSE_031_2026.04_Rev.01.docx",
    "checklist": BASE / "HMRC_Checklist RM_031_2026.04_Rev.01.xlsx",
    "incident_board": BASE / "HMRC_Quadro de Sinistralidade_RM_031_2026.04.xlsx",
    "formacoes": BASE / "Formações_compressed.pdf",
    "toolbox": BASE / "TOOLBOX_compressed (2).pdf",
    "quadros": BASE / "Inspenções de Quadro.pdf",
    "equipamento": BASE / "Inspenção equipamento.pdf",
    "inducao": BASE / "Indução 4.pdf",
    "extintor": BASE / "Extintor.pdf",
    "andaime": BASE / "Andaime.pdf",
    "advertencias": BASE / "Advertencias.pdf",
}

OUT = Path("extracted-cabinda")


def norm(value):
    if value is None:
        return ""
    text = str(value).replace("\xa0", " ").strip()
    return " ".join(text.split())


def extract_docx(path: Path) -> dict:
    document = Document(path)
    paragraphs = [norm(p.text) for p in document.paragraphs if norm(p.text)]
    tables = []
    for table_index, table in enumerate(document.tables, start=1):
        rows = []
        for row in table.rows:
            values = [norm(cell.text) for cell in row.cells]
            if any(values):
                rows.append(values)
        tables.append({"table": table_index, "rows": rows})
    return {"path": str(path), "paragraphs": paragraphs, "tables": tables}


def extract_xlsx(path: Path) -> dict:
    workbook = load_workbook(path, data_only=False)
    sheets = []
    for sheet in workbook.worksheets:
        rows = []
        for row in sheet.iter_rows():
            values = [norm(cell.value) for cell in row]
            if any(values):
                rows.append({"row": row[0].row, "values": values})
        sheets.append(
            {
                "title": sheet.title,
                "max_row": sheet.max_row,
                "max_column": sheet.max_column,
                "rows": rows,
            }
        )
    return {"path": str(path), "sheets": sheets}


def extract_pdf(path: Path) -> dict:
    info = {"path": str(path), "pages": None, "text": []}
    if PdfReader is None:
        return info
    reader = PdfReader(str(path))
    info["pages"] = len(reader.pages)
    for index, page in enumerate(reader.pages, start=1):
        text = norm(page.extract_text() or "")
        if text:
            info["text"].append({"page": index, "text": text})
    return info


def write_doc_markdown(name: str, data: dict) -> None:
    lines = [f"# {name}", "", f"Source: `{data['path']}`", ""]
    if data.get("paragraphs"):
        lines.extend(["## Paragraphs", ""])
        for i, paragraph in enumerate(data["paragraphs"], start=1):
            lines.append(f"{i}. {paragraph}")
        lines.append("")
    for table in data.get("tables", []):
        lines.extend([f"## Table {table['table']}", ""])
        for row in table["rows"]:
            lines.append(" | ".join(row))
        lines.append("")
    (OUT / f"{name}.md").write_text("\n".join(lines), encoding="utf-8")


def write_xlsx_markdown(name: str, data: dict) -> None:
    lines = [f"# {name}", "", f"Source: `{data['path']}`", ""]
    for sheet in data["sheets"]:
        lines.extend([f"## Sheet: {sheet['title']} ({sheet['max_row']} x {sheet['max_column']})", ""])
        for row in sheet["rows"]:
            lines.append(f"R{row['row']}: " + " | ".join(row["values"]))
        lines.append("")
    (OUT / f"{name}.md").write_text("\n".join(lines), encoding="utf-8")


def write_pdf_markdown(name: str, data: dict) -> None:
    lines = [f"# {name}", "", f"Source: `{data['path']}`", "", f"Pages: {data['pages']}", ""]
    for page in data.get("text", []):
        lines.extend([f"## Page {page['page']}", "", page["text"], ""])
    (OUT / f"{name}.md").write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    OUT.mkdir(exist_ok=True)
    manifest = {}
    for name, path in FILES.items():
        if not path.exists():
            manifest[name] = {"missing": str(path)}
            continue
        suffix = path.suffix.lower()
        if suffix == ".docx":
            data = extract_docx(path)
            write_doc_markdown(name, data)
            manifest[name] = {
                "paragraphs": len(data["paragraphs"]),
                "tables": len(data["tables"]),
                "table_rows": [len(table["rows"]) for table in data["tables"]],
            }
        elif suffix == ".xlsx":
            data = extract_xlsx(path)
            write_xlsx_markdown(name, data)
            manifest[name] = {
                "sheets": [
                    {
                        "title": sheet["title"],
                        "max_row": sheet["max_row"],
                        "max_column": sheet["max_column"],
                        "non_empty_rows": len(sheet["rows"]),
                    }
                    for sheet in data["sheets"]
                ]
            }
        elif suffix == ".pdf":
            data = extract_pdf(path)
            write_pdf_markdown(name, data)
            manifest[name] = {"pages": data["pages"], "text_pages": len(data.get("text", []))}
    print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
